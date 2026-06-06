"""Rebuild Kimberly's resume with professional formatting — preserves all content."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"c:\Users\morph\OneDrive\MyResume-Professional.docx"

# Healthcare-friendly palette
PRIMARY = RGBColor(0x1B, 0x49, 0x65)      # deep teal-navy
ACCENT = RGBColor(0x2E, 0x86, 0xAB)       # bright teal
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x2D, 0x34, 0x36)
TEXT_MUTED = RGBColor(0x55, 0x65, 0x73)

SKILLS = [
    "Hospice & Palliative Care",
    "Strong Attention to Detail",
    "Resident Support",
    "Dementia Care",
    "First Aid and CPR",
    "ADL Management",
    "Multitasking Abilities",
]

SUMMARY = (
    "Healthcare professional with strong foundation in resident care and support. "
    "Recognized for delivering high-quality care and effectively managing residents' daily routines. "
    "Valued for reliability and ability to work collaboratively with team members to adapt to changing "
    "needs and ensure consistent results. Proficient in assisting with daily activities, maintaining "
    "safe environment, and providing emotional support."
)

JOB1_BULLETS = [
    "Assisted residents in preparing for activity and social programs.",
    "Maintained clean, safe, and well-organized patient environment.",
    "Attended, participated, and contributed to monthly staff meetings addressing resident needs.",
    "Responded to room transfers, incident reports, and maintenance requests.",
    "Assisted residents with daily dental and mouth care, bath functions, and hair care.",
    "Enhanced resident satisfaction by promptly addressing concerns and resolving conflicts within the community.",
    "Changed bed linens, dumped trash, and smoothly handled maintenance issues to promote resident comfort.",
    "Provided emotional support for residents during difficult transitions or personal challenges, enhancing overall mental wellbeing.",
    "Maintained a clean, organized living space for each resident, fostering a comfortable and welcoming atmosphere.",
    "Answered call lights and supported patient comfort and safety by adjusting bed rails and equipment.",
    "Received recognition from supervisors for consistently delivering high-quality care under challenging circumstances.",
    "Contributed to efficient facility operations by maintaining accurate records of resident information and care provided.",
    "Improved residents' well-being through consistent monitoring and timely assistance with personal hygiene tasks.",
    "Ensured compliance with state regulations regarding resident rights, confidentiality requirements, and facility procedures.",
    "Developed strong relationships with family members, providing regular updates on resident progress and addressing concerns as needed.",
    "Contributed to the development of individualized care plans, ensuring accurate documentation of each resident's needs.",
    "Supported residents during meal times, assisting with feeding and promoting proper nutrition.",
    "Enhanced resident satisfaction by providing compassionate and attentive care.",
    "Furthered skills by actively taking part in employee training and taking classes to improve skills.",
    "Followed residence's service plan and provided assistance with toileting, bathing, dressing, oral hygiene and other daily living tasks.",
    "Delivered and served meals and provided other dining needs.",
    "Escorted residents to community activities and games and encouraged participation.",
    "Conversed with assigned residents to determine emotional and physical state of being and reported changes to resident services director.",
    "Completed activities of daily living for patients unable to self-care, and assisted those with limited mobility in completing tasks.",
    "Cared for clients with diagnoses such as respiratory failure, diabetes, Parkinson's disease and muscular dystrophy.",
    "Responded to patient requests for supplies and personal comfort items such as extra blankets.",
    "Facilitated personal hygiene management, feeding and ambulation.",
    "Helped patients with self-feeding and assisted feeding, based on individual needs.",
    "Changed linens in rooms, keeping spaces fresh and clean for patient health and satisfaction.",
    "Observed patients under care conditions to help identify symptoms, responses to treatments and progress with goals.",
]

JOB2_BULLETS = [
    "Provided personalized daily care to clients, including mobility assistance, hygiene support, and medication administration.",
    "Assisted with physical therapy routines and safe patient transfers using proper body mechanics.",
    "Monitored patient conditions, documenting observations and reporting significant changes to supervisors.",
    "Offered emotional support and companionship, contributing to the overall well-being of clients and their families.",
    "Assisted clients with laundry, light housekeeping, and other tasks to maintain a clean living environment.",
    "Monitored client health status and reported significant changes to healthcare professionals for timely interventions.",
    "Supported clients with personal grooming, including bathing, dressing, and oral hygiene, to promote dignity and self-esteem.",
    "Encouraged independence by assisting clients with mobility and physical exercises tailored to individual capabilities.",
    "Monitored changes in clients' conditions to report concerns to the supervisor.",
    "Supported families through difficult times by offering emotional support and education on important care tasks.",
    "Assisted with end-of-life care.",
    "Communicated regularly with clients' families to provide updates on health and well-being.",
    "Scheduled and coordinated medical appointments.",
    "Administered medications in accordance with the doctor's instructions.",
    "Monitored client vital signs, administered medications, and tracked behaviors to keep the healthcare supervisor well-informed.",
    "Assisted with feeding and monitored intake to help patients achieve nutritional objectives.",
    "Recognized and reported abnormalities or changes in patients' health status to the case manager.",
    "Turned and positioned bedbound patients to prevent bedsores and maintain comfort levels.",
    "Built strong relationships with clients to deliver emotional support and companionship.",
    "Engaged patients in meaningful conversation, socialization, and activity while providing personal care assistance.",
    "Assisted with dressing guidance, grooming, meal preparation, and medication reminders.",
    "Changed dressings, bandages, and binders to maintain proper healing and sanitary measures.",
    "Maintained clean personal areas and prepared healthy meals to support the client's nutritional needs.",
    "Promoted overall well-being by preparing nutritious meals according to individual dietary needs and preferences.",
    "Assisted patients with self-administered medications.",
]

REFERENCES = [
    "Lesha Caldon: (214) 529-6785",
    "Steve Rossini: (412) 999-7375",
    "Gerald Webb: (682) 559-6862",
    "Jeff Hall: (214) 406-2356",
]

# Unified palette — one family throughout
BG_PANEL = "E8F4F8"
BG_PANEL_ALT = "F4F9FB"
BAR_ACCENT = "2E86AB"
BAR_DARK = "1B4965"


def set_cell_shading(cell, color_hex):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def remove_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, val in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def add_colored_header(doc):
    table = doc.add_table(rows=1, cols=1)
    remove_table_borders(table)
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "1B4965")
    set_cell_margins(cell, top=200, bottom=200, left=200, right=200)

    name_p = cell.paragraphs[0]
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_after = Pt(2)
    nr = name_p.add_run("Kimberly Mosher")
    nr.bold = True
    nr.font.size = Pt(26)
    nr.font.color.rgb = WHITE
    nr.font.name = "Calibri"

    tagline_p = cell.add_paragraph()
    tagline_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tagline_p.paragraph_format.space_after = Pt(6)
    tr = tagline_p.add_run("Healthcare & Resident Care Professional")
    tr.font.size = Pt(11)
    tr.font.color.rgb = RGBColor(0xA8, 0xD8, 0xEA)
    tr.font.italic = True

    accent_p = cell.add_paragraph()
    accent_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    accent_p.paragraph_format.space_after = Pt(8)
    ar = accent_p.add_run("━" * 42)
    ar.font.size = Pt(8)
    ar.font.color.rgb = ACCENT

    contact_p = cell.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_after = Pt(2)
    cr = contact_p.add_run("Irving, TX 75060   •   (214) 277-1238   •   dakm5908@yahoo.com")
    cr.font.size = Pt(10)
    cr.font.color.rgb = WHITE

    profile_p = cell.add_paragraph()
    profile_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pr = profile_p.add_run("Bold Profile")
    pr.font.size = Pt(10)
    pr.font.color.rgb = RGBColor(0xC8, 0xE6, 0xF0)
    pr.font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_section_heading(doc, text):
    table = doc.add_table(rows=1, cols=2)
    remove_table_borders(table)
    table.autofit = False
    table.columns[0].width = Inches(0.08)
    table.columns[1].width = Inches(6.42)

    bar = table.cell(0, 0)
    set_cell_shading(bar, BAR_ACCENT)
    set_cell_margins(bar, top=40, bottom=40, left=0, right=0)

    body = table.cell(0, 1)
    set_cell_shading(body, BG_PANEL)
    set_cell_margins(body, top=60, bottom=60, left=140, right=100)
    p = body.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = PRIMARY
    run.font.name = "Calibri"

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def add_panel_row(doc, lines, bold_first=False):
    """Single cohesive panel with teal left bar — used for summary, education, cert."""
    table = doc.add_table(rows=1, cols=2)
    remove_table_borders(table)
    table.autofit = False
    table.columns[0].width = Inches(0.08)
    table.columns[1].width = Inches(6.42)

    bar = table.cell(0, 0)
    set_cell_shading(bar, BAR_ACCENT)
    set_cell_margins(bar, top=40, bottom=40, left=0, right=0)

    body = table.cell(0, 1)
    set_cell_shading(body, BG_PANEL_ALT)
    set_cell_margins(body, top=100, bottom=100, left=140, right=120)

    for i, line in enumerate(lines):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.font.size = Pt(10.5)
        run.font.color.rgb = TEXT if not bold_first or i > 0 else PRIMARY
        run.bold = bold_first and i == 0
        run.font.name = "Calibri"

    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def add_skill_list(doc):
    table = doc.add_table(rows=1, cols=2)
    remove_table_borders(table)
    table.autofit = False
    table.columns[0].width = Inches(0.08)
    table.columns[1].width = Inches(6.42)

    bar = table.cell(0, 0)
    set_cell_shading(bar, BAR_ACCENT)

    body = table.cell(0, 1)
    set_cell_shading(body, BG_PANEL_ALT)
    set_cell_margins(body, top=100, bottom=100, left=140, right=120)

    inner = body.add_table(rows=(len(SKILLS) + 1) // 2, cols=2)
    remove_table_borders(inner)
    for i, skill in enumerate(SKILLS):
        cell = inner.cell(i // 2, i % 2)
        set_cell_margins(cell, top=40, bottom=40, left=0, right=80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        marker = p.add_run("▸ ")
        marker.font.color.rgb = ACCENT
        marker.font.size = Pt(10.5)
        run = p.add_run(skill)
        run.font.size = Pt(10.5)
        run.font.color.rgb = TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def add_job_header(doc, title, dates, company):
    table = doc.add_table(rows=2, cols=1)
    remove_table_borders(table)

    title_cell = table.cell(0, 0)
    set_cell_shading(title_cell, "1B4965")
    set_cell_margins(title_cell, top=80, bottom=40, left=160, right=160)
    tp = title_cell.paragraphs[0]
    tp.paragraph_format.space_after = Pt(0)
    tr = tp.add_run(title)
    tr.bold = True
    tr.font.size = Pt(12)
    tr.font.color.rgb = WHITE
    sep = tp.add_run("   •   ")
    sep.font.color.rgb = ACCENT
    dr = tp.add_run(dates)
    dr.font.size = Pt(10.5)
    dr.font.color.rgb = RGBColor(0xA8, 0xD8, 0xEA)
    dr.font.italic = True

    co_cell = table.cell(1, 0)
    set_cell_shading(co_cell, BG_PANEL)
    set_cell_margins(co_cell, top=40, bottom=80, left=160, right=160)
    cp = co_cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(0)
    cr = cp.add_run(company)
    cr.font.size = Pt(10.5)
    cr.font.color.rgb = PRIMARY
    cr.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.15)
        marker = p.add_run("▸ ")
        marker.font.color.rgb = ACCENT
        marker.font.size = Pt(10.5)
        body = p.add_run(item)
        body.font.size = Pt(10.5)
        body.font.color.rgb = TEXT


def add_references_list(doc):
    table = doc.add_table(rows=1, cols=2)
    remove_table_borders(table)
    table.autofit = False
    table.columns[0].width = Inches(0.08)
    table.columns[1].width = Inches(6.42)

    bar = table.cell(0, 0)
    set_cell_shading(bar, BAR_ACCENT)

    body = table.cell(0, 1)
    set_cell_shading(body, BG_PANEL_ALT)
    set_cell_margins(body, top=100, bottom=100, left=140, right=120)

    inner = body.add_table(rows=2, cols=2)
    remove_table_borders(inner)
    for i, ref in enumerate(REFERENCES):
        row, col = divmod(i, 2)
        cell = inner.cell(row, col)
        set_cell_margins(cell, top=40, bottom=40, left=0, right=80)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        name, phone = ref.split(": ")
        nr = p.add_run(name)
        nr.bold = True
        nr.font.size = Pt(10.5)
        nr.font.color.rgb = PRIMARY
        p.add_run("\n")
        pr = p.add_run(phone)
        pr.font.size = Pt(10)
        pr.font.color.rgb = TEXT_MUTED


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT

    add_colored_header(doc)

    add_section_heading(doc, "Professional Summary")
    add_panel_row(doc, [SUMMARY])

    add_section_heading(doc, "Skills")
    add_skill_list(doc)

    add_section_heading(doc, "Work History")

    add_job_header(
        doc,
        "Resident Assistant",
        "August 2025 – Present",
        "Washington Pointe – Las Colinas, Texas",
    )
    add_bullets(doc, JOB1_BULLETS)

    doc.add_paragraph().paragraph_format.space_before = Pt(8)

    add_job_header(
        doc,
        "In-Home Care Provider",
        "September 2012 – June 2022",
        "Girling Health Care – Dallas, TX",
    )
    add_bullets(doc, JOB2_BULLETS)

    add_section_heading(doc, "Education")
    add_panel_row(
        doc,
        ["High School Diploma, Honors Academy", "Irving, TX  •  May 2000"],
        bold_first=True,
    )

    add_section_heading(doc, "Certifications")
    add_panel_row(doc, ["BLS (Basic Life Support) & CPR Certified"])

    add_section_heading(doc, "References")
    add_references_list(doc)

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
